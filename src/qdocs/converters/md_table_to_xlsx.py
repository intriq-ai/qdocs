"""Extract GFM markdown tables and export to XLSX."""

from __future__ import annotations

from pathlib import Path

from qdocs.exceptions import ConversionError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    _AVAILABLE = True
except ImportError:
    _AVAILABLE = False


def _split_row(line: str) -> list[str]:
    inner = line.strip().strip("|")
    out: list[str] = []
    token: list[str] = []
    escaped = False
    for char in inner:
        if escaped:
            token.append(char)
            escaped = False
            continue
        if char == "\\":
            escaped = True
            token.append(char)
            continue
        if char == "|":
            out.append("".join(token).strip())
            token = []
            continue
        token.append(char)
    out.append("".join(token).strip())
    return out


def _is_separator(line: str) -> bool:
    row = _split_row(line)
    return bool(row) and all(c.replace(":", "").replace("-", "").strip() == "" for c in row)


def _extract_tables(text: str) -> list[tuple[str, list[str], list[list[str]]]]:
    """Return list of (context_name, header, rows)."""
    tables: list[tuple[str, list[str], list[list[str]]]] = []
    current_context = "Data"

    lines = text.splitlines()
    idx = 0
    while idx < len(lines):
        stripped = lines[idx].strip()

        if stripped.startswith("## "):
            current_context = stripped[3:].strip() or "Data"
            idx += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            header = _split_row(stripped)
            if idx + 1 >= len(lines):
                idx += 1
                continue
            sep = lines[idx + 1].strip()
            if not (sep.startswith("|") and sep.endswith("|") and _is_separator(sep)):
                idx += 1
                continue

            rows: list[list[str]] = []
            idx += 2
            while idx < len(lines):
                row_line = lines[idx].strip()
                if not (row_line.startswith("|") and row_line.endswith("|")):
                    break
                rows.append(_split_row(row_line))
                idx += 1

            tables.append((current_context, header, rows))
            continue

        idx += 1

    return tables


def convert_md_tables_to_xlsx(
    source: Path,
    target: Path,
    *,
    one_table_per_sheet: bool = True,
) -> None:
    """Extract markdown tables and write an XLSX workbook."""
    if not _AVAILABLE:
        raise ConversionError("openpyxl not installed. Run: uv pip install openpyxl")
    if not source.exists():
        raise ConversionError(f"Source file not found: {source}")

    text = source.read_text(encoding="utf-8")
    tables = _extract_tables(text)
    if not tables:
        raise ConversionError("No markdown tables found in source")

    workbook = openpyxl.Workbook()
    if "Sheet" in workbook.sheetnames:
        del workbook["Sheet"]

    header_fill = PatternFill(fill_type="solid", fgColor="D6E4F0")
    header_font = Font(bold=True)

    if not one_table_per_sheet:
        ws = workbook.create_sheet(title="Tables")
        current_row = 1
        for idx, (context, headers, rows) in enumerate(tables, start=1):
            ws.cell(row=current_row, column=1, value=f"Table {idx}: {context}").font = Font(bold=True)
            current_row += 1
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True)
            current_row += 1
            for row_values in rows:
                for col_idx, value in enumerate(row_values, start=1):
                    ws.cell(row=current_row, column=col_idx, value=value).alignment = Alignment(wrap_text=True)
                current_row += 1
            current_row += 2
    else:
        for idx, (context, headers, rows) in enumerate(tables, start=1):
            sheet_name = f"{idx:02d}-{context}"[:31]
            ws = workbook.create_sheet(title=sheet_name)
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(header) + 4, 10), 60)
            for row_idx, row_values in enumerate(rows, start=2):
                for col_idx, value in enumerate(row_values, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=value).alignment = Alignment(wrap_text=True)
            ws.freeze_panes = ws["A2"]

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(target))
