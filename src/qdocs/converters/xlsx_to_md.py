"""XLSX → Markdown converter using openpyxl."""

from pathlib import Path

from loguru import logger

from qdocs.exceptions import ConversionError

try:
    import openpyxl

    _AVAILABLE = True
except ImportError:
    _AVAILABLE = False


def convert_xlsx_to_md(source: Path, target: Path) -> None:
    """Convert an XLSX workbook to Markdown tables.

    Raises:
        ConversionError: If openpyxl is unavailable or conversion fails.
    """
    if not _AVAILABLE:
        msg = "openpyxl not installed. Run: uv pip install openpyxl"
        raise ConversionError(msg)
    if not source.exists():
        msg = f"Source not found: {source}"
        raise ConversionError(msg)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.load_workbook(source, data_only=True)
        lines: list[str] = [f"# {source.stem}\n"]

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            if len(workbook.sheetnames) > 1:
                lines.append(f"\n## {sheet_name}\n")
            header = [str(c) if c is not None else "" for c in rows[0]]
            lines.append(f"| {' | '.join(header)} |\n")
            lines.append(f"| {' | '.join(['---'] * len(header))} |\n")
            for row in rows[1:]:
                cells = [str(c) if c is not None else "" for c in row]
                lines.append(f"| {' | '.join(cells)} |\n")

        workbook.close()
        target.write_text("\n".join(lines), encoding="utf-8")
        logger.debug(f"XLSX→MD: {source.name} → {target.name}")
    except Exception as exc:
        msg = f"Failed to convert {source.name} to MD: {exc}"
        raise ConversionError(msg) from exc
