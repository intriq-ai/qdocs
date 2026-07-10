"""Pre-flight syntax validators for qdocs conversion workflows."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from qdocs.exceptions import ConversionError

try:
    import openpyxl

    _XLSX_AVAILABLE = True
except ImportError:
    _XLSX_AVAILABLE = False


@dataclass(slots=True)
class ValidationIssue:
    """Structured validation issue with location context."""

    message: str
    line: int | None = None
    column: int | None = None


@dataclass(slots=True)
class ValidationResult:
    """Validation output with errors/warnings split."""

    kind: Literal["csv", "md", "xlsx"]
    errors: list[ValidationIssue]
    warnings: list[ValidationIssue]

    @property
    def is_valid(self) -> bool:
        return not self.errors


_MD_TABLE_LINE = re.compile(r"^\|.*\|$")
_MD_TABLE_SEPARATOR = re.compile(r"^\|\s*[:\-][:\-\s|]*\|$")


def _split_md_row(line: str) -> list[str]:
    """Split a markdown pipe row while preserving escaped pipes."""
    inner = line.strip().strip("|")
    cells: list[str] = []
    token: list[str] = []
    escaped = False
    for char in inner:
        if escaped:
            token.append(char)
            escaped = False
            continue
        if char == "\\":
            escaped = True
            token.append(char)
            continue
        if char == "|":
            cells.append("".join(token).strip())
            token = []
            continue
        token.append(char)
    cells.append("".join(token).strip())
    return cells


def validate_csv(source: Path, *, encoding: str = "utf-8") -> ValidationResult:
    """Validate CSV structure and delimiter consistency.

    Checks include:
    - Parseability via csv module
    - Detectable dialect (with fallback)
    - Column-count consistency across rows
    """
    if not source.exists():
        raise ConversionError(f"Source file not found: {source}")

    errors: list[ValidationIssue] = []
    warnings: list[ValidationIssue] = []

    text = source.read_text(encoding=encoding)
    if not text.strip():
        warnings.append(ValidationIssue("CSV file is empty"))
        return ValidationResult(kind="csv", errors=errors, warnings=warnings)

    sample = text[:8192]
    delimiter = ","
    try:
        dialect = csv.Sniffer().sniff(sample)
        delimiter = dialect.delimiter
    except csv.Error:
        warnings.append(ValidationIssue("Could not infer CSV dialect; using comma delimiter fallback"))

    expected_columns: int | None = None
    try:
        reader = csv.reader(text.splitlines(), delimiter=delimiter)
        for idx, row in enumerate(reader, start=1):
            if expected_columns is None:
                expected_columns = len(row)
                continue
            if len(row) != expected_columns:
                warnings.append(
                    ValidationIssue(
                        f"Row has {len(row)} columns; expected {expected_columns}",
                        line=idx,
                    )
                )
    except csv.Error as exc:
        errors.append(ValidationIssue(f"CSV parse error: {exc}"))

    return ValidationResult(kind="csv", errors=errors, warnings=warnings)


def validate_md_tables(source: Path, *, require_separator: bool = True) -> ValidationResult:
    """Validate markdown table syntax (GFM pipe tables)."""
    if not source.exists():
        raise ConversionError(f"Source file not found: {source}")

    errors: list[ValidationIssue] = []
    warnings: list[ValidationIssue] = []

    lines = source.read_text(encoding="utf-8").splitlines()
    in_table = False
    header_cols = 0
    saw_separator = False

    for idx, raw in enumerate(lines, start=1):
        line = raw.rstrip()

        if _MD_TABLE_LINE.match(line):
            cells = _split_md_row(line)
            if not in_table:
                in_table = True
                header_cols = len(cells)
                saw_separator = False
                if header_cols == 0:
                    errors.append(ValidationIssue("Table header row has no cells", line=idx))
                continue

            if _MD_TABLE_SEPARATOR.match(line):
                saw_separator = True
                sep_cols = len(cells)
                if sep_cols != header_cols:
                    errors.append(
                        ValidationIssue(
                            f"Separator columns ({sep_cols}) do not match header columns ({header_cols})",
                            line=idx,
                        )
                    )
                continue

            if len(cells) != header_cols:
                warnings.append(
                    ValidationIssue(
                        f"Data row has {len(cells)} columns; header has {header_cols}",
                        line=idx,
                    )
                )
            continue

        if in_table:
            if require_separator and not saw_separator:
                errors.append(ValidationIssue("Table is missing separator row (| --- |)", line=idx - 1))
            in_table = False
            header_cols = 0
            saw_separator = False

    if in_table and require_separator and not saw_separator:
        errors.append(ValidationIssue("Table is missing separator row (| --- |)", line=len(lines)))

    return ValidationResult(kind="md", errors=errors, warnings=warnings)


def validate_xlsx(source: Path) -> ValidationResult:
    """Validate basic XLSX workbook integrity and worksheet metadata."""
    if not _XLSX_AVAILABLE:
        raise ConversionError("openpyxl not installed. Run: uv pip install openpyxl")
    if not source.exists():
        raise ConversionError(f"Source file not found: {source}")

    errors: list[ValidationIssue] = []
    warnings: list[ValidationIssue] = []

    try:
        wb = openpyxl.load_workbook(source, data_only=False)
        seen_lower: set[str] = set()
        for sheet in wb.worksheets:
            title = sheet.title
            if len(title) > 31:
                errors.append(ValidationIssue("Sheet title exceeds 31 characters"))
            lower = title.lower()
            if lower in seen_lower:
                errors.append(ValidationIssue(f"Duplicate sheet title (case-insensitive): {title}"))
            seen_lower.add(lower)
            if sheet.max_row == 1 and sheet.max_column == 1 and sheet.cell(1, 1).value is None:
                warnings.append(ValidationIssue(f"Sheet '{title}' is empty"))
        wb.close()
    except Exception as exc:
        errors.append(ValidationIssue(f"Failed to open workbook: {exc}"))

    return ValidationResult(kind="xlsx", errors=errors, warnings=warnings)
