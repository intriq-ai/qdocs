"""CSV -> XLSX converter with type inference and basic styling."""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path
from typing import Any

from qdocs.exceptions import ConversionError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    _AVAILABLE = True
except ImportError:
    _AVAILABLE = False


def _looks_like_int(value: str) -> bool:
    try:
        int(value.replace(",", ""))
        return True
    except ValueError:
        return False


def _looks_like_float(value: str) -> bool:
    try:
        float(value.replace(",", ""))
        return True
    except ValueError:
        return False


def _looks_like_date(value: str) -> datetime | None:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def _infer_value(raw: str) -> tuple[Any, str | None]:
    value = raw.strip()
    if value == "":
        return None, None

    if value.startswith("$"):
        clean = value[1:].replace(",", "").strip()
        try:
            return float(clean), '"$"#,##0.00'
        except ValueError:
            return value, None

    parsed_date = _looks_like_date(value)
    if parsed_date:
        return parsed_date, "YYYY-MM-DD"

    if _looks_like_int(value):
        return int(value.replace(",", "")), "#,##0"

    if _looks_like_float(value):
        return float(value.replace(",", "")), "#,##0.00"

    return value, None


def _detect_dialect(text: str) -> csv.Dialect:
    sample = text[:8192]
    try:
        return csv.Sniffer().sniff(sample)
    except csv.Error:
        return csv.get_dialect("excel")


def convert_csv_to_xlsx(
    source: Path,
    target: Path,
    *,
    sheet_name: str = "Data",
    delimiter: str | None = None,
    has_header: bool = True,
    auto_width: bool = True,
) -> None:
    """Convert CSV file into XLSX workbook with inferred cell formats."""
    if not _AVAILABLE:
        raise ConversionError("openpyxl not installed. Run: uv pip install openpyxl")
    if not source.exists():
        raise ConversionError(f"Source file not found: {source}")

    text = source.read_text(encoding="utf-8")
    if not text.strip():
        raise ConversionError("CSV source is empty")

    dialect = _detect_dialect(text)
    if delimiter:
        dialect.delimiter = delimiter

    try:
        rows = list(csv.reader(text.splitlines(), dialect=dialect))
    except csv.Error as exc:
        raise ConversionError(f"CSV parse error: {exc}") from exc

    if not rows:
        raise ConversionError("CSV source is empty")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]

    header_fill = PatternFill(fill_type="solid", fgColor="D6E4F0")
    header_font = Font(bold=True)

    col_widths: dict[int, int] = {}

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, raw in enumerate(row, start=1):
            if has_header and row_idx == 1:
                cell = sheet.cell(row=row_idx, column=col_idx, value=raw)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                value, number_format = _infer_value(raw)
                cell = sheet.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                if number_format:
                    cell.number_format = number_format

            width = len(str(raw)) if raw else 0
            col_widths[col_idx] = max(col_widths.get(col_idx, 8), min(width + 2, 80))

    if has_header:
        sheet.freeze_panes = sheet["A2"]

    if auto_width:
        for col_idx, width in col_widths.items():
            sheet.column_dimensions[get_column_letter(col_idx)].width = width

    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(target))
