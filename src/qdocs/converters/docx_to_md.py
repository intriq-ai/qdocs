"""DOCX → Markdown converter."""

from pathlib import Path

from docx import Document
from loguru import logger

from qdocs.exceptions import ConversionError


def convert_docx_to_md(source: Path, target: Path) -> None:
    """Convert a DOCX file to Markdown.

    Raises:
        ConversionError: On failure.
    """
    if not source.exists():
        msg = f"Source not found: {source}"
        raise ConversionError(msg)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        doc = Document(str(source))
        lines: list[str] = [f"# {source.stem}\n"]

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style else ""
            if style.startswith("Heading "):
                level = style.replace("Heading ", "")
                prefix = "#" * int(level) if level.isdigit() else ""
                lines.append(f"{prefix} {text}\n" if prefix else f"{text}\n")
            elif style == "List Bullet":
                lines.append(f"- {text}\n")
            elif style == "List Number":
                lines.append(f"1. {text}\n")
            else:
                lines.append(f"{text}\n")

        for table in doc.tables:
            lines.append("\n")
            for i, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                lines.append(f"| {' | '.join(cells)} |\n")
                if i == 0:
                    lines.append(f"| {' | '.join(['---'] * len(cells))} |\n")

        target.write_text("\n".join(lines), encoding="utf-8")
        logger.debug(f"DOCX→MD: {source.name} → {target.name}")
    except Exception as exc:
        msg = f"Failed to convert {source.name} to MD: {exc}"
        raise ConversionError(msg) from exc
