"""PPTX → Markdown converter using python-pptx."""

from pathlib import Path

from loguru import logger

from qdocs.exceptions import ConversionError

try:
    from pptx import Presentation

    _AVAILABLE = True
except ImportError:
    _AVAILABLE = False


def convert_pptx_to_md(source: Path, target: Path) -> None:
    """Convert a PPTX presentation to Markdown.

    Raises:
        ConversionError: If python-pptx is unavailable or conversion fails.
    """
    if not _AVAILABLE:
        msg = "python-pptx not installed. Run: uv pip install python-pptx"
        raise ConversionError(msg)
    if not source.exists():
        msg = f"Source not found: {source}"
        raise ConversionError(msg)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        prs = Presentation(str(source))
        lines: list[str] = [f"# {source.stem}\n"]

        for slide_num, slide in enumerate(prs.slides, 1):
            lines.append(f"\n## Slide {slide_num}\n")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    lines.append(f"{shape.text.strip()}\n")
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        lines.append(f"| {' | '.join(cells)} |\n")

        target.write_text("\n".join(lines), encoding="utf-8")
        logger.debug(f"PPTX→MD: {source.name} → {target.name}")
    except Exception as exc:
        msg = f"Failed to convert {source.name} to MD: {exc}"
        raise ConversionError(msg) from exc
