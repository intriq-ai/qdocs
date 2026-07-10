"""PDF → DOCX converter using pdf2docx."""

import logging
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from loguru import logger
from pdf2docx import Converter

from qdocs.exceptions import ConversionError


def _remove_trailing_blank_paragraphs(docx_path: Path) -> None:
    try:
        doc = Document(str(docx_path))
        to_remove = []
        for para in reversed(doc.paragraphs):
            if para.text.strip():
                break
            to_remove.append(para)
        for para in to_remove:
            para._element.getparent().remove(para._element)
        doc.save(str(docx_path))
    except Exception as exc:
        logger.warning(f"Blank page cleanup failed for {docx_path.name}: {exc}")


def _apply_landscape(docx_path: Path) -> None:
    try:
        doc = Document(str(docx_path))
        for section in doc.sections:
            pw = int(section.page_width or 0)  # type: ignore[arg-type]
            ph = int(section.page_height or 0)  # type: ignore[arg-type]
            if pw > ph:
                pg_sz = section._sectPr.xpath("./w:pgSz")
                if pg_sz:
                    pg_sz[0].set(qn("w:orient"), "landscape")
        doc.save(str(docx_path))
    except Exception as exc:
        logger.warning(f"Landscape orientation failed for {docx_path.name}: {exc}")


def convert_pdf_to_docx(
    source: Path, target: Path, landscape_mode: bool = False
) -> None:
    """Convert a PDF to DOCX via pdf2docx.

    Args:
        source: Source PDF file.
        target: Target DOCX file.
        landscape_mode: Apply landscape orientation after conversion.

    Raises:
        ConversionError: On failure.
    """
    if not source.exists():
        msg = f"Source not found: {source}"
        raise ConversionError(msg)
    try:
        target.parent.mkdir(parents=True, exist_ok=True)

        # Suppress pdf2docx verbose logging
        for name in ("pdf2docx", ""):
            lg = logging.getLogger(name)
            orig = (lg.level, lg.propagate)
            lg.setLevel(logging.CRITICAL)
            lg.propagate = False

        cv = Converter(str(source))
        cv.convert(str(target))
        cv.close()

        for name in ("pdf2docx", ""):
            lg = logging.getLogger(name)
            lg.setLevel(orig[0])
            lg.propagate = orig[1]

        _remove_trailing_blank_paragraphs(target)
        if landscape_mode:
            _apply_landscape(target)

        logger.debug(f"PDF→DOCX: {source.name} → {target.name}")
    except Exception as exc:
        msg = f"Failed to convert {source.name} to DOCX: {exc}"
        raise ConversionError(msg) from exc
